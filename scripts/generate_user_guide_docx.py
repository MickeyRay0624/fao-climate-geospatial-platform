#!/usr/bin/env python3
"""Generate the polished English platform user guide from its Markdown source.

Design system:
  - selected preset: compact_reference_guide
  - first-page header template: editorial_cover
  - named overrides used consistently:
      compact_table_text: 9pt for 2-column tables, 8.4pt for 3+ columns
      cover_identity: editorial-cover colours and 30pt title
      callout_panel: preset callout fill with a heading-blue left rule
"""

from __future__ import annotations

import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/user-guide/Climate_Geospatial_DSS_User_Guide_2026-09-02.md"
OUTPUT = ROOT / "docs/user-guide/Climate_Geospatial_DSS_User_Guide_2026-09-02.docx"

FONT = "Calibri"
HEADING_BLUE = "2E74B5"
HEADING_DARK_BLUE = "1F4D78"
INK_BLUE = "0B2545"
TABLE_FILL = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
POSITIVE_DARK_BLUE = "1F3A5F"
CAUTION_GOLD = "7A5A00"
RISK_RED = "9B1C1C"
MUTED = "596675"
RULE_GRAY = "D7DBE2"
WHITE = "FFFFFF"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
MARGIN_DXA = 1440
HEADER_FOOTER_DXA = 708
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP_BOTTOM_DXA = 80
CELL_MARGIN_START_END_DXA = 120


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, name: str = FONT) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), name)


def set_widow_control(paragraph, enabled: bool = True) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    element = ppr.find(qn("w:widowControl"))
    if element is None:
        element = OxmlElement("w:widowControl")
        ppr.append(element)
    element.set(qn("w:val"), "1" if enabled else "0")


def set_cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top: int, start: int, bottom: int, end: int) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcmar = tcpr.find(qn("w:tcMar"))
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcmar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tcmar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(width_dxa))
    tcw.set(qn("w:type"), "dxa")
    cell.width = Inches(width_dxa / 1440)


def set_table_geometry(table, widths: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}, got {sum(widths)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = table._tbl.tblPr

    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tblw.set(qn("w:type"), "dxa")

    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(indent_dxa))
    tblind.set(qn("w:type"), "dxa")

    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = borders.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:space"), "0")
        item.set(qn("w:color"), "B8C4D1")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        cant_split = trpr.find(qn("w:cantSplit"))
        if cant_split is None:
            trpr.append(OxmlElement("w:cantSplit"))
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(
                cell,
                CELL_MARGIN_TOP_BOTTOM_DXA,
                CELL_MARGIN_START_END_DXA,
                CELL_MARGIN_TOP_BOTTOM_DXA,
                CELL_MARGIN_START_END_DXA,
            )
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    header = trpr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        trpr.append(header)
    header.set(qn("w:val"), "true")


def add_paragraph_rule(paragraph, color: str = RULE_GRAY, size: str = "6") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = pbdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        pbdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, result, end])


def add_hyperlink(paragraph, text: str, url: str, color: str = HEADING_BLUE) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)
    rpr.append(rfonts)
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), color)
    rpr.append(colour)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    new_run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, *, size: float | None = None, color: str | None = None) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run)
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = rgb(color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run)
            run.bold = True
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = rgb(color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Courier New")
            run.font.size = Pt((size or 11) - 0.8)
            run.font.color.rgb = rgb(INK_BLUE)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), LIGHT_GRAY)
            run._r.get_or_add_rPr().append(shading)
        elif token.startswith("["):
            link_match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link_match:
                add_hyperlink(paragraph, link_match.group(1), link_match.group(2), color or HEADING_BLUE)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run)
            run.italic = True
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = rgb(color)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run)
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = rgb(color)


def create_numbering(document: Document) -> tuple[int, int, int]:
    numbering = document.part.numbering_part.element
    existing_abstract = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abstract = max(existing_abstract or [0]) + 1
    next_num = max(existing_num or [0]) + 1

    def add_definition(abstract_id: int, num_id: int, fmt: str, level_text: str) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        numfmt = OxmlElement("w:numFmt")
        numfmt.set(qn("w:val"), fmt)
        level.append(numfmt)
        lvltext = OxmlElement("w:lvlText")
        lvltext.set(qn("w:val"), level_text)
        level.append(lvltext)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        level.append(lvl_jc)
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ppr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "271")
        ppr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.append(spacing)
        level.append(ppr)
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{key}"), FONT)
        rpr.append(rfonts)
        level.append(rpr)
        abstract.append(level)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)

    bullet_abstract = next_abstract
    bullet_num = next_num
    add_definition(bullet_abstract, bullet_num, "bullet", "•")
    number_abstract = next_abstract + 1
    number_num = next_num + 1
    add_definition(number_abstract, number_num, "decimal", "%1.")
    return bullet_num, number_num, number_abstract


def create_numbering_instance(document: Document, abstract_id: int, start_at: int = 1) -> int:
    """Create a fresh list instance so each independent procedure restarts at 1."""
    numbering = document.part.numbering_part.element
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    num_id = max(existing_num or [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start_at))
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.append(numpr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, num_id_element])


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK_BLUE)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_values = {
        "Heading 1": (16, HEADING_BLUE, 18, 10),
        "Heading 2": (13, HEADING_BLUE, 14, 7),
        "Heading 3": (12, HEADING_DARK_BLUE, 10, 5),
    }
    for style_name, (size, colour, before, after) in heading_values.items():
        style = styles[style_name]
        style.base_style = normal
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(colour)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True


def add_header_content(header, left: str, right: str) -> None:
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(4)
    tabs = paragraph.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.5), alignment=2)
    run = paragraph.add_run(left)
    set_run_font(run)
    run.bold = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = rgb(MUTED)
    run = paragraph.add_run("\t" + right)
    set_run_font(run)
    run.font.size = Pt(8.5)
    run.font.color.rgb = rgb(MUTED)
    add_paragraph_rule(paragraph, RULE_GRAY, "6")


def add_footer_content(footer, left: str) -> None:
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(4)
    tabs = paragraph.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.5), alignment=2)
    add_paragraph_rule(paragraph, RULE_GRAY, "4")
    run = paragraph.add_run(left)
    set_run_font(run)
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(MUTED)
    run = paragraph.add_run("\tPage ")
    set_run_font(run)
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(MUTED)
    add_field(run, "PAGE")
    run = paragraph.add_run(" of ")
    set_run_font(run)
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(MUTED)
    add_field(run, "NUMPAGES")


def configure_headers_and_footers(section) -> None:
    add_header_content(
        section.first_page_header,
        "CLIMATE GEOSPATIAL",
        "HOSTED TEAM REVIEW EDITION",
    )
    add_footer_content(section.first_page_footer, "dss.lmqstudio.com · Demonstration evidence only")
    add_header_content(
        section.header,
        "Climate Geospatial Data & Decision Platform",
        "User Guide · v0.3.0",
    )
    add_footer_content(section.footer, "dss.lmqstudio.com · 2 September 2026")


def add_spacer(document: Document, points: float) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(points)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run("")
    set_run_font(run)


def add_cover(document: Document) -> None:
    add_spacer(document, 132)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("PLATFORM USER GUIDE")
    set_run_font(run)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = rgb(CAUTION_GOLD)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Climate Geospatial\nData & Decision Platform")
    set_run_font(run)
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = rgb("203748")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("User Guide, Capability Overview & Product Roadmap")
    set_run_font(run)
    run.font.size = Pt(15)
    run.font.color.rgb = rgb("2B5163")

    edition = document.add_paragraph()
    edition.alignment = WD_ALIGN_PARAGRAPH.CENTER
    edition.paragraph_format.space_after = Pt(88)
    run = edition.add_run("—  CAMBODIA RICE RESILIENCE · HOSTED TEAM REVIEW  —")
    set_run_font(run)
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = rgb(CAUTION_GOLD)

    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_after = Pt(4)
    run = date.add_run("2 September 2026 · Version 0.3.0-demo-platform")
    set_run_font(run)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = rgb("203748")

    audience = document.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    audience.paragraph_format.space_after = Pt(10)
    run = audience.add_run("For programme leaders, data teams, analysts, extension teams and technical evaluators")
    set_run_font(run)
    run.italic = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = rgb("505050")

    url = document.add_paragraph()
    url.alignment = WD_ALIGN_PARAGRAPH.CENTER
    url.paragraph_format.space_after = Pt(12)
    add_hyperlink(url, "https://dss.lmqstudio.com", "https://dss.lmqstudio.com", HEADING_BLUE)

    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run("EVIDENCE BOUNDARY")
    set_run_font(run)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = rgb(HEADING_DARK_BLUE)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    add_inline(
        paragraph,
        "The hosted edition uses synthetic analysis data and fictional field cases. It proves workflow and engineering capability; it is not operational advice, a funding recommendation or an agronomic diagnosis.",
        size=9.2,
        color=INK_BLUE,
    )


def table_widths(column_count: int) -> list[int]:
    if column_count == 1:
        return [9360]
    if column_count == 2:
        return [2700, 6660]
    if column_count == 3:
        return [2200, 3300, 3860]
    if column_count == 4:
        return [1800, 2520, 2520, 2520]
    base = CONTENT_WIDTH_DXA // column_count
    widths = [base] * column_count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_markdown_table(document: Document, raw_lines: list[str]) -> None:
    rows = [parse_table_row(line) for line in raw_lines]
    if len(rows) > 1 and is_separator_row(rows[1]):
        rows.pop(1)
    column_count = len(rows[0])
    if any(len(row) != column_count for row in rows):
        raise ValueError(f"Inconsistent Markdown table near: {raw_lines[0]}")
    widths = table_widths(column_count)
    table = document.add_table(rows=len(rows), cols=column_count)
    set_table_geometry(table, widths)
    repeat_table_header(table.rows[0])
    font_size = 9.0 if column_count <= 2 else 8.4

    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            if row_index == 0:
                set_cell_shading(cell, TABLE_FILL)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.05
            set_widow_control(paragraph)
            add_inline(paragraph, value, size=font_size, color=INK_BLUE)
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = rgb(HEADING_DARK_BLUE)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_callout(document: Document, lines: list[str]) -> None:
    text_lines = [line.strip() for line in lines if line.strip()]
    if not text_lines:
        return
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), HEADING_BLUE)
    borders.append(left)
    tcpr.append(borders)

    for index, line in enumerate(text_lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2 if index == 0 else 0)
        paragraph.paragraph_format.line_spacing = 1.15
        add_inline(paragraph, line, size=9.4, color=INK_BLUE)
        if index == 0 and line.startswith("**"):
            for run in paragraph.runs:
                run.font.color.rgb = rgb(HEADING_DARK_BLUE)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_body_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    set_widow_control(paragraph)
    add_inline(paragraph, text)


def add_list_paragraph(document: Document, text: str, num_id: int) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    set_widow_control(paragraph)
    apply_numbering(paragraph, num_id)
    add_inline(paragraph, text)


def render_markdown(
    document: Document,
    source: str,
    bullet_num: int,
    decimal_num: int,
    decimal_abstract: int,
) -> None:
    lines = source.splitlines()
    first_break = lines.index("<!-- PAGE BREAK -->")
    lines = lines[first_break + 1 :]
    index = 0
    paragraph_buffer: list[str] = []
    active_decimal_num: int | None = None
    first_decimal_sequence = True

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = " ".join(item.strip() for item in paragraph_buffer).strip()
            if text:
                add_body_paragraph(document, text)
            paragraph_buffer = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if stripped == "<!-- PAGE BREAK -->":
            flush_paragraph()
            active_decimal_num = None
            # The Markdown marker is a source-reading aid. In the DOCX, natural
            # pagination avoids stranding a continuation table on an otherwise
            # empty page; Heading styles keep the next paragraph with the title.
            index += 1
            continue

        if stripped == "---":
            flush_paragraph()
            active_decimal_num = None
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(8)
            add_paragraph_rule(paragraph, RULE_GRAY, "8")
            index += 1
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            active_decimal_num = None
            level = len(heading_match.group(1))
            paragraph = document.add_paragraph(style=f"Heading {level}")
            set_widow_control(paragraph)
            add_inline(paragraph, heading_match.group(2))
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            active_decimal_num = None
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(document, table_lines)
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            active_decimal_num = None
            callout_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                callout_lines.append(lines[index].strip()[1:].strip())
                index += 1
            add_callout(document, callout_lines)
            continue

        bullet_match = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet_match:
            flush_paragraph()
            active_decimal_num = None
            add_list_paragraph(document, bullet_match.group(1), bullet_num)
            index += 1
            continue

        number_match = re.match(r"^\d+\.\s+(.+)$", stripped)
        if number_match:
            flush_paragraph()
            if active_decimal_num is None:
                if first_decimal_sequence:
                    active_decimal_num = decimal_num
                    first_decimal_sequence = False
                else:
                    active_decimal_num = create_numbering_instance(document, decimal_abstract)
            add_list_paragraph(document, number_match.group(1), active_decimal_num)
            index += 1
            continue

        active_decimal_num = None
        paragraph_buffer.append(stripped.rstrip("  "))
        index += 1

    flush_paragraph()


def set_update_fields(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def audit_output(path: Path) -> None:
    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        styles_xml = archive.read("word/styles.xml").decode("utf-8")
        numbering_xml = archive.read("word/numbering.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")

    requirements = {
        "Letter page width": 'w:pgSz w:w="12240"',
        "Letter page height": 'w:h="15840"',
        "1-inch margins": 'w:top="1440"',
        "header distance": 'w:header="708"',
        "footer distance": 'w:footer="708"',
        "fixed table width": 'w:tblW w:type="dxa" w:w="9360"',
        "table indent": 'w:tblInd w:w="120" w:type="dxa"',
        "fixed table layout": 'w:tblLayout w:type="fixed"',
        "update fields": 'w:updateFields w:val="true"',
    }
    failures: list[str] = []
    for label, needle in requirements.items():
        haystack = settings_xml if label == "update fields" else document_xml
        if needle not in haystack:
            failures.append(label)

    style_requirements = [
        'w:styleId="Normal"',
        'w:after="120"',
        'w:line="300"',
        'w:styleId="Heading1"',
        'w:before="360"',
        'w:after="200"',
        'w:styleId="Heading2"',
        'w:before="280"',
        'w:after="140"',
        'w:styleId="Heading3"',
        'w:before="200"',
        'w:after="100"',
    ]
    for needle in style_requirements:
        if needle not in styles_xml:
            failures.append(f"style token {needle}")

    numbering_requirements = [
        'w:pos="540"',
        'w:left="540"',
        'w:hanging="271"',
        'w:after="80"',
        'w:line="300"',
        'w:numFmt w:val="bullet"',
        'w:numFmt w:val="decimal"',
    ]
    for needle in numbering_requirements:
        if needle not in numbering_xml:
            failures.append(f"numbering token {needle}")

    if failures:
        raise RuntimeError("DOCX design-token audit failed: " + ", ".join(failures))


def main() -> int:
    source = SOURCE.read_text(encoding="utf-8")
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    configure_section(section)
    configure_headers_and_footers(section)
    bullet_num, decimal_num, decimal_abstract = create_numbering(document)
    set_update_fields(document)

    properties = document.core_properties
    properties.title = "Climate Geospatial Data & Decision Platform — User Guide"
    properties.subject = "Capability overview, user instructions and product roadmap"
    properties.author = "LMQ Studio"
    properties.keywords = "geospatial, decision support, data governance, investment prioritisation, extension"
    properties.comments = "Hosted team-review edition; demonstration evidence only."

    add_cover(document)
    document.add_page_break()
    render_markdown(document, source, bullet_num, decimal_num, decimal_abstract)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    audit_output(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Size: {OUTPUT.stat().st_size:,} bytes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
